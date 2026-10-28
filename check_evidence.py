import openpyxl

wb = openpyxl.load_workbook('StarWars_QA_Bug_Report.xlsx')
ws = wb['Defect Log with Evidence']

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
evidence_col = None
for idx, h in enumerate(headers, 1):
    if 'Evidence' in str(h) or 'Proof' in str(h):
        evidence_col = idx
        break

print(f"Evidence Column: {evidence_col} ({headers[evidence_col-1]})")
print("=" * 80)

missing_count = 0
for r in range(2, ws.max_row + 1):
    num = ws.cell(row=r, column=1).value
    bid = ws.cell(row=r, column=2).value
    title = ws.cell(row=r, column=3).value
    evid = ws.cell(row=r, column=evidence_col).value
    
    if not evid or len(str(evid).strip()) == 0:
        print(f"[MISSING] Row {r}: {bid} has NO evidence!")
        missing_count += 1
    else:
        # format preview
        lines = str(evid).strip().split('\n')
        print(f"{num:02d}. [{bid}] {title}")
        print(f"    Evidence: {lines[0]}")
        if len(lines) > 1:
            print(f"    Details:  {lines[1][:90]}...")
        print("-" * 80)

print(f"\nTotal Bugs Checked: {ws.max_row - 1}")
print(f"Bugs with Evidence: {ws.max_row - 1 - missing_count}")
print(f"Bugs Missing Evidence: {missing_count}")
